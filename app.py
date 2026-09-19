import streamlit as st
import pandas as pd
import io
import calendar
import json
import base64
from streamlit_local_storage import LocalStorage
import streamlit.components.v1 as components
import openpyxl

st.set_page_config(page_title="ระบบสรุปยอดทำงานลูกจ้าง (แบบ 51)", layout="wide")

# ==========================================
# 1. ฟังก์ชัน LocalStorage
# ==========================================
local_storage = LocalStorage()

def save_roster_to_local(df):
    if df is not None and not df.empty:
        roster_json = df.to_json(orient='records')
        roster_b64 = base64.b64encode(roster_json.encode('utf-8')).decode('utf-8')
        js_code = f"""
        <script>
            window.parent.localStorage.setItem('srt_form51_data', atob('{roster_b64}'));
        </script>
        """
        components.html(js_code, height=0, width=0)

# ==========================================
# 2. ข้อมูลตั้งต้น
# ==========================================
months_list = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", 
               "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
shift_codes = ["", "/", "พ", "ป", "น", "ย", "2น", "2ย"]

st.title("📝 ระบบสรุปยอดทำงานและวันหยุดลูกจ้าง (Export แบบ 51)")

# ==========================================
# 3. ตั้งค่าข้อมูลส่วนกลาง (รอบเดือน & ตัวแปร)
# ==========================================
with st.container(border=True):
    st.subheader("⚙️ 1. ตั้งค่าข้อมูลส่วนกลาง")
    
    st.markdown("##### 📅 รอบเดือนที่เบิก และฐานข้อมูล")
    c1, c2, c3 = st.columns([1, 1, 2])
    with c1:
        target_month_name = st.selectbox("เดือนปัจจุบัน", months_list, index=7)
    with c2:
        target_year_be = st.number_input("ปี พ.ศ.", value=2569)
    with c3:
        uploaded_db = st.file_uploader("📂 อัปโหลดไฟล์ Excel ฐานข้อมูลพนักงาน", type=["xlsx"])
        
    target_month_idx = months_list.index(target_month_name) + 1
    if target_month_idx == 1:
        prev_month_idx = 12
        prev_year_be = target_year_be - 1
    else:
        prev_month_idx = target_month_idx - 1
        prev_year_be = target_year_be
        
    prev_month_name = months_list[prev_month_idx - 1]
    _, num_days_prev = calendar.monthrange(prev_year_be - 543, prev_month_idx)

    st.markdown("---")
    st.markdown("##### 📝 ข้อมูลตัวแปร (โปรแกรมจะนำไปแทนที่ตัวแปรในฟอร์มแบบอัตโนมัติ)")
    
    # เพิ่มช่อง [DATE] กลับมาให้แล้วครับ
    col_contract, col_doc_date = st.columns(2)
    with col_contract:
        contract_no = st.text_input("เลขที่สัญญา/คำสั่ง [CONTRACT]", value="")
    with col_doc_date:
        doc_date = st.text_input("วันที่เซ็นแบบฟอร์ม [DATE]", value="")
        
    col_date1, col_date2, col_date3 = st.columns(3)
    with col_date1:
        date_contract = st.text_input("วันที่ลงนามสัญญา [DATE_CONTRACT]", value="")
    with col_date2:
        sdate_contract = st.text_input("วันที่เริ่มสัญญา [SDATE_CONTRACT]", value="")
    with col_date3:
        edate_contract = st.text_input("วันสิ้นสุดสัญญา [EDATE_CONTRACT]", value="")

# ==========================================
# 4. เตรียมคอลัมน์ตารางหน้าเว็บ (อัปเดตระบบพิมพ์เร็ว)
# ==========================================
# ล็อค 3 คอลัมน์แรกให้อยู่ติดขอบซ้ายเสมอ
config = {
    "ชื่อ-สกุล": st.column_config.TextColumn("ชื่อ-สกุล", width="medium", pinned=True),
    "เลขประจำตัว": st.column_config.TextColumn("เลขประจำตัว", width="small", pinned=True),
    "อัตราวันละ": st.column_config.NumberColumn("อัตราวันละ", width="small", pinned=True)
}

columns_list = ["ชื่อ-สกุล", "เลขประจำตัว", "อัตราวันละ"]

# เปลี่ยนคอลัมน์วันที่ให้เป็น TextColumn แทน SelectboxColumn เพื่อให้คีย์บอร์ดเลื่อนเร็วได้
prev_cols = [f"P{d}" for d in range(1, num_days_prev + 1)]
for d, col in enumerate(prev_cols, 1):
    columns_list.append(col)
    config[col] = st.column_config.TextColumn(f"{d} {prev_month_name[:3]}.", width="small")

curr_cols = [f"C{d}" for d in range(1, 16)]
for d, col in enumerate(curr_cols, 1):
    columns_list.append(col)
    config[col] = st.column_config.TextColumn(f"{d} {target_month_name[:3]}.", width="small")

# ==========================================
# 5. โหลดข้อมูล (จากไฟล์อัปโหลด หรือ LocalStorage)
# ==========================================
saved_roster_json = local_storage.getItem("srt_form51_data")

# 🛡️ แก้ไขบั๊ก F5: ป้องกันการสร้างตารางเปล่ามาทับก่อน LocalStorage โหลดเสร็จ
if 'loaded_from_ls' not in st.session_state:
    st.session_state.loaded_from_ls = False

if saved_roster_json and not st.session_state.loaded_from_ls:
    try:
        st.session_state.form51_data = pd.read_json(io.StringIO(saved_roster_json), orient='records')
        st.session_state.loaded_from_ls = True
    except:
        pass

if 'form51_data' not in st.session_state:
    st.session_state.form51_data = pd.DataFrame(columns=columns_list)

if uploaded_db is not None:
    try:
        df_db = pd.read_excel(uploaded_db, sheet_name=0)
        name_col = next((col for col in ["รายชื่อ", "ชื่อ-สกุล", "ชื่อ-นามสกุล", "ชื่อ"] if col in df_db.columns), None)
        id_col = next((col for col in ["เลขประจำตัว", "รหัสพนักงาน", "ID"] if col in df_db.columns), None)
        rate_col = next((col for col in ["อัตราวันละ", "ค่าแรง", "ค่าจ้าง", "อัตรา"] if col in df_db.columns), None)
        
        if name_col:
            new_df = pd.DataFrame(columns=columns_list)
            new_df["ชื่อ-สกุล"] = df_db[name_col].astype(str)
            if id_col: new_df["เลขประจำตัว"] = df_db[id_col].astype(str)
            if rate_col: new_df["อัตราวันละ"] = pd.to_numeric(df_db[rate_col], errors='coerce')

            new_df = new_df.fillna("").replace("nan", "") 
            
            st.session_state.form51_data = new_df
            save_roster_to_local(new_df)
            st.success("✅ โหลดรายชื่อ พร้อมเลขประจำตัวและอัตราค่าจ้างสำเร็จ!")
    except Exception as e:
        st.error(f"อ่านไฟล์ฐานข้อมูลไม่สำเร็จ: {e}")

# ==========================================
# 6. แสดงตารางกรอกข้อมูล
# ==========================================
# 🎯 ถ้ามีการกดปุ่มกู้คืน หรือยกยอดมาจากด้านล่าง ให้ทำการเซฟลงเครื่องที่จุดนี้ทันที
if st.session_state.get('pending_save', False):
    save_roster_to_local(st.session_state.form51_data)
    st.session_state.pending_save = False

st.markdown("### ✍️ 2. ตารางกรอกข้อมูลลงเวลา (46 วัน)")
with st.form("editor_form"):
    edited_df = st.data_editor(
        st.session_state.form51_data,
        num_rows="dynamic",
        column_config=config,
        use_container_width=True,
        height=500
    )
    submit_btn = st.form_submit_button("💾 บันทึกข้อมูลลงเครื่องเบราว์เซอร์ (กดบ่อยๆ กันเหนียว)", type="secondary")
    if submit_btn:
        st.session_state.form51_data = edited_df
        save_roster_to_local(edited_df)
        st.success("บันทึกข้อมูลไว้ในเบราว์เซอร์เรียบร้อยแล้ว!")

# ==========================================
# 🛡️ ระบบ Backup และ ยกยอดไปเดือนถัดไป
# ==========================================
st.markdown("---")
st.markdown("### 🛡️ ระบบจัดการไฟล์ Backup & เริ่มเดือนใหม่")
c_back1, c_back2 = st.columns(2)

with c_back1:
    st.info("💡 **เซฟงานเก็บไว้:** ดาวน์โหลดข้อมูลที่กรอกไว้เป็นไฟล์ Excel")
    
    backup_output = io.BytesIO()
    with pd.ExcelWriter(backup_output, engine='xlsxwriter') as writer:
        edited_df.to_excel(writer, index=False, sheet_name='Backup')
    backup_output.seek(0)
    
    st.download_button(
        label="📥 ดาวน์โหลดไฟล์ Backup",
        data=backup_output,
        file_name=f"Backup_แบบ51_{target_month_name}_{target_year_be}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )

with c_back2:
    st.info("🔄 **นำไฟล์มาใช้:** อัปโหลดไฟล์ Backup ของเดือนนี้ หรือ 'เดือนที่แล้ว'")
    uploaded_backup = st.file_uploader("📂 อัปโหลดไฟล์ Backup ของคุณที่นี่", type=["xlsx"], key="backup_upload")
    
    if uploaded_backup is not None:
        try:
            df_backup = pd.read_excel(uploaded_backup)
            if "ชื่อ-สกุล" in df_backup.columns and "เลขประจำตัว" in df_backup.columns:
                df_backup = df_backup.fillna("").astype(str).replace("nan", "")
                
                c_btn1, c_btn2 = st.columns(2)
                
                with c_btn1:
                    if st.button("✨ กู้คืน (ทำเดือนเดิมต่อ)", use_container_width=True, type="primary"):
                        st.session_state.form51_data = df_backup
                        st.session_state.pending_save = True # สั่งให้เซฟในรอบถัดไป
                        st.rerun() # 🚀 รีเฟรชแอปตัวเองทันที! (ไม่ต้องกด F5)
                        
                with c_btn2:
                    if st.button("⏭️ ยกยอด (เริ่มเดือนใหม่)", use_container_width=True, type="secondary"):
                        new_df = pd.DataFrame(columns=columns_list)
                        new_df["ชื่อ-สกุล"] = df_backup["ชื่อ-สกุล"]
                        new_df["เลขประจำตัว"] = df_backup["เลขประจำตัว"]
                        new_df["อัตราวันละ"] = df_backup.get("อัตราวันละ", "")
                        
                        for i in range(1, 16):
                            old_c = f"C{i}"
                            new_p = f"P{i}"
                            if old_c in df_backup.columns and new_p in new_df.columns:
                                new_df[new_p] = df_backup[old_c]
                                
                        new_df = new_df.fillna("")
                        st.session_state.form51_data = new_df
                        st.session_state.pending_save = True # สั่งให้เซฟในรอบถัดไป
                        st.rerun() # 🚀 รีเฟรชแอปตัวเองทันที! (ไม่ต้องกด F5)
            else:
                st.warning("⚠️ ไฟล์นี้ไม่ใช่ไฟล์ Backup แบบฟอร์ม 51 ครับ")
        except Exception as e:
            st.error(f"ไฟล์ Backup ไม่ถูกต้อง: {e}")
# ==========================================
# 7. ระบบคำนวณและ Export 
# ==========================================
st.markdown("---")
if st.button("📊 คำนวณและส่งออกไฟล์ Excel (ฟอร์ม 51)", type="primary", use_container_width=True):
    if edited_df.empty or edited_df["ชื่อ-สกุล"].isnull().all():
        st.warning("กรุณากรอกชื่อพนักงานอย่างน้อย 1 คน")
    else:
        # ========================================================
        # 🛠️ ระบบจัดเรียงข้อมูลใหม่ตาม "เลขประจำตัว" (น้อยไปมาก)
        # ========================================================
        sort_df = edited_df.copy()
        # สร้างคอลัมน์จำลองเพื่อแปลงเลขประจำตัวเป็นตัวเลข (ป้องกันปัญหา 10 มาก่อน 2)
        sort_df['sort_key'] = pd.to_numeric(sort_df['เลขประจำตัว'], errors='coerce')
        # สั่งเรียงลำดับจากน้อยไปมาก (คนที่ไม่มีเลขประจำตัวจะถูกดันไปไว้ล่างสุด)
        sort_df = sort_df.sort_values(by='sort_key', ascending=True, na_position='last').drop(columns=['sort_key'])
        
        work_data = []    
        holiday_data = [] 

        def map_work_sheet(v):
            v = str(v).strip()
            if v == "2น": return "น"
            if v == "2ย": return "ย"
            if v in ["/", "พ", "ป", "น", "ย"]: return v
            return ""

        def map_holiday_sheet(v):
            v = str(v).strip()
            if v == "2น": return "2/น"  
            if v == "2ย": return "2/ย"  
            return "" 
        
        # 🔄 เปลี่ยนมาวนลูปอ่านข้อมูลจากตารางที่จัดเรียงแล้ว (sort_df)
        for idx, row in sort_df.iterrows():
            name = str(row.get("ชื่อ-สกุล", "")).strip()
            if not name or name == "nan": continue
            
            # เก็บค่าเลขประจำตัวและอัตรา ไว้โยนลง Excel
            emp_id = str(row.get("เลขประจำตัว", "")).strip()
            emp_rate = str(row.get("อัตราวันละ", "")).strip()
            if emp_id == "nan": emp_id = ""
            if emp_rate == "nan": emp_rate = ""
            
            curr_1_15_raw = [str(row.get(c, "")).strip() for c in curr_cols]
            prev_1_15_raw = [str(row.get(prev_cols[i], "")).strip() for i in range(0, 15)]
            prev_16_end_raw = [str(row.get(prev_cols[i], "")).strip() for i in range(15, num_days_prev)]
            
            # 📁 ชีท "ค่าทำงาน"
            work_curr_1_15 = [map_work_sheet(v) for v in curr_1_15_raw]
            work_prev_16_end = [map_work_sheet(v) for v in prev_16_end_raw]
            working_range = work_curr_1_15 + work_prev_16_end
            
            count_normal = working_range.count("/")
            count_n = working_range.count("น")
            count_vac = working_range.count("พ")
            count_sick = working_range.count("ป")
            
            row_work = {
                "ที่": len(work_data) + 1,  # ลำดับที่ จะรันใหม่ 1, 2, 3... ตามที่เรียงสวยๆ แล้ว
                "ชื่อ-นามสกุล": name,
                "เลขประจำตัว": emp_id,
                "อัตราวันละ": emp_rate
            }
            
            for i in range(1, 16): row_work[str(i)] = work_curr_1_15[i-1]
            for i in range(16, num_days_prev + 1): row_work[str(i)] = work_prev_16_end[i-16]
            row_work["ปกติ"] = count_normal
            row_work["พรบ."] = count_n + count_vac + count_sick
            row_work["373"] = count_normal
            row_work[".1 (ป)"] = count_sick
            row_work[".2 (พ)"] = count_vac
            row_work[".6 (น)"] = count_n
            work_data.append(row_work)
            
            # 🏖️ ชีท "วันหยุด"
            holiday_prev_1_15 = [map_holiday_sheet(v) for v in prev_1_15_raw]
            holiday_prev_16_end = [map_holiday_sheet(v) for v in prev_16_end_raw]
            holiday_range = holiday_prev_1_15 + holiday_prev_16_end
            
            count_2n = holiday_range.count("2/น")
            count_2y = holiday_range.count("2/ย")
            val_6 = count_2n * 1
            val_7 = count_2y * 2
            
            row_holiday = {
                "ที่": len(holiday_data) + 1, # ลำดับที่ รันใหม่เช่นเดียวกัน
                "ชื่อ-นามสกุล": name,
                "เลขประจำตัว": emp_id,
                "อัตราวันละ": emp_rate
            }
            
            for i in range(1, 16): row_holiday[str(i)] = holiday_prev_1_15[i-1]
            for i in range(16, num_days_prev + 1): row_holiday[str(i)] = holiday_prev_16_end[i-16]
            row_holiday["พรบ."] = val_6 + val_7
            row_holiday[".6 (2/น)"] = val_6
            row_holiday[".7 (2/ย)"] = val_7
            holiday_data.append(row_holiday)

        chunk_size = 13
        work_chunks = [work_data[i:i + chunk_size] for i in range(0, len(work_data), chunk_size)]
        holiday_chunks = [holiday_data[i:i + chunk_size] for i in range(0, len(holiday_data), chunk_size)]

        # ==========================================
        # 🖨️ สร้างไฟล์ Excel ลงฟอร์มต้นแบบ
        # ==========================================
        import openpyxl
        output = io.BytesIO()
        try:
            wb = openpyxl.load_workbook("template_51.xlsx")
            ws_work_template = wb["ค่าทำงาน"]  
            ws_holiday_template = wb["วันหยุด"] 

            def replace_tags_in_sheet(ws, page_num):
                replacements = {
                    "[MONTH]": f"{target_month_name} {target_year_be}",
                    "[PMONTH]": str(prev_month_name),       
                    "[N]": str(num_days_prev),              
                    "[PAGE]": str(page_num),                
                    "[CONTRACT]": str(contract_no),
                    "[DATE]": str(doc_date),                # <--- เพิ่ม [DATE] กลับเข้าสู่ระบบแล้ว!
                    "[DATE_CONTRACT]": str(date_contract),
                    "[SDATE_CONTRACT]": str(sdate_contract),
                    "[EDATE_CONTRACT]": str(edate_contract)
                }
                
                for r in range(1, 60):
                    for c in range(1, 45):
                        cell = ws.cell(row=r, column=c)
                        if cell.value and isinstance(cell.value, str):
                            for tag, actual_value in replacements.items():
                                if tag in cell.value:
                                    cell.value = cell.value.replace(tag, str(actual_value))

            # ----------------------------------------
            # 1. จัดการฝั่ง "ค่าทำงาน"
            # ----------------------------------------
            for page_idx, chunk in enumerate(work_chunks):
                ws = wb.copy_worksheet(ws_work_template)
                ws.title = f"ค่าทำงาน_หน้า{page_idx + 1}"
                
                replace_tags_in_sheet(ws, page_num=page_idx + 1)
                
                start_row = 6 
                for i, person_data in enumerate(chunk):
                    current_row = start_row + i
                    ws.cell(row=current_row, column=1).value = person_data["ที่"]
                    ws.cell(row=current_row, column=2).value = person_data["ชื่อ-นามสกุล"]
                    ws.cell(row=current_row, column=3).value = person_data["เลขประจำตัว"] 
                    ws.cell(row=current_row, column=4).value = person_data["อัตราวันละ"]  
                    
                    for d in range(1, 32):
                        col_idx = 4 + d  
                        ws.cell(row=current_row, column=col_idx).value = person_data.get(str(d), "")
                        
                    ws.cell(row=current_row, column=36).value = person_data["ปกติ"] 
                    ws.cell(row=current_row, column=37).value = person_data["พรบ."] 
                    ws.cell(row=current_row, column=38).value = person_data["373"]  
                    ws.cell(row=current_row, column=39).value = person_data[".1 (ป)"]
                    ws.cell(row=current_row, column=40).value = person_data[".2 (พ)"]
                    ws.cell(row=current_row, column=41).value = person_data[".6 (น)"]

            # ----------------------------------------
            # 2. จัดการฝั่ง "วันหยุด"
            # ----------------------------------------
            for page_idx, chunk in enumerate(holiday_chunks):
                ws = wb.copy_worksheet(ws_holiday_template)
                ws.title = f"วันหยุด_หน้า{page_idx + 1}"
                
                replace_tags_in_sheet(ws, page_num=page_idx + 1)
                
                start_row = 6 
                for i, person_data in enumerate(chunk):
                    current_row = start_row + i
                    ws.cell(row=current_row, column=1).value = person_data["ที่"]
                    ws.cell(row=current_row, column=2).value = person_data["ชื่อ-นามสกุล"]
                    ws.cell(row=current_row, column=3).value = person_data["เลขประจำตัว"] 
                    ws.cell(row=current_row, column=4).value = person_data["อัตราวันละ"]  
                    
                    for d in range(1, 32):
                        col_idx = 4 + d 
                        ws.cell(row=current_row, column=col_idx).value = person_data.get(str(d), "")
                        
                    ws.cell(row=current_row, column=37).value = person_data["พรบ."] 
                    ws.cell(row=current_row, column=43).value = person_data[".6 (2/น)"] 
                    ws.cell(row=current_row, column=45).value = person_data[".7 (2/ย)"] 

            wb.remove(ws_work_template)
            wb.remove(ws_holiday_template)
            wb.save(output)
            output.seek(0)
            
            st.success(f"✅ คำนวณเสร็จสมบูรณ์! และจัดเรียงรายชื่อตาม 'เลขประจำตัว' จากน้อยไปมากให้เรียบร้อยแล้ว!")
            st.download_button(
                label="📥 ดาวน์โหลดไฟล์ฟอร์ม 51 (พร้อมปริ้นท์)",
                data=output,
                file_name=f"ฟอร์ม51_{target_month_name}_{target_year_be}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"เกิดข้อผิดพลาด: {e}")
